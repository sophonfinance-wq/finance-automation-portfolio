"""Synthetic, fully-fictional sample workpaper generator.

Everything produced here is fake by construction: entity names come from a
hard-coded list of obviously-invented placeholders ("Demo Holdings LLC",
"Maple Fund LP", ...), and all figures are drawn from the stdlib :mod:`random`
module seeded for full reproducibility. No real entity, person, EIN, path or
dollar amount appears anywhere.

The generator can emit two flavours of workpaper:

* ``clean``    -- internally consistent; every formula ties out.
* ``defective`` -- seeded with a handful of realistic defects (a tie-out
  break, a hard-coded value that should be a formula, an unsupported
  AI-assumption, and leaked internal/process language) so the pipeline has
  something to catch.

It can also write the workpaper to a tiny ``.xlsx`` via :mod:`openpyxl` so the
package is self-contained for an Excel-shaped demo.
"""

from __future__ import annotations

import random
from typing import List, Tuple

from triangulate.model import AuthoritySource, Workpaper, WorkpaperCell

# Obviously-fake entity placeholders -- never a real company.
_FAKE_ENTITIES: List[Tuple[str, str]] = [
    ("Demo Holdings LLC", "FY24-Q4"),
    ("Maple Fund LP", "FY24-Q4"),
    ("Birchwood Op Co", "FY24-Q4"),
    ("Cedar Path Ventures Inc", "FY24-Q4"),
    ("Acme Sandbox Trust", "FY24-Q4"),
]

# Leaked "process language" the audit must catch never reaching a client cell.
_LEAKED_PROCESS_PHRASES = [
    "TODO: ask the LLM to recheck this",
    "as the AI suggested, assume 12%",
    "placeholder until Reviewer confirms",
]


def _new_workpaper(rng: random.Random) -> Tuple[Workpaper, float, float, float]:
    """Create the common skeleton and return key driver figures."""
    entity, period = rng.choice(_FAKE_ENTITIES)
    wp = Workpaper(
        engagement="ENG-DEMO-0001",
        entity=entity,
        period=period,
    )

    # Three fictional revenue streams (round, obviously-synthetic numbers).
    rev_a = float(rng.randrange(100_000, 400_000, 1_000))
    rev_b = float(rng.randrange(50_000, 200_000, 1_000))
    rev_c = float(rng.randrange(10_000, 90_000, 1_000))
    return wp, rev_a, rev_b, rev_c


def make_clean_workpaper(seed: int = 20240101) -> Workpaper:
    """Build an internally consistent workpaper where every formula ties out."""
    rng = random.Random(seed)
    wp, rev_a, rev_b, rev_c = _new_workpaper(rng)

    total_rev = rev_a + rev_b + rev_c
    tax_rate = 0.21  # flat, documented assumption
    tax = round(total_rev * tax_rate, 2)
    net = round(total_rev - tax, 2)

    wp.set_cell(WorkpaperCell("B2", "Revenue - Stream A", rev_a,
                              source=AuthoritySource.CURRENT_YEAR_SOURCE))
    wp.set_cell(WorkpaperCell("B3", "Revenue - Stream B", rev_b,
                              source=AuthoritySource.CURRENT_YEAR_SOURCE))
    wp.set_cell(WorkpaperCell("B4", "Revenue - Stream C", rev_c,
                              source=AuthoritySource.CURRENT_YEAR_SOURCE))
    wp.set_cell(WorkpaperCell("B5", "Total Revenue", total_rev, formula="=B2+B3+B4",
                              source=AuthoritySource.WORKBOOK_FORMULA))
    wp.set_cell(WorkpaperCell("B6", "Tax Rate (documented)", tax_rate,
                              source=AuthoritySource.MANAGEMENT_INSTRUCTION))
    wp.set_cell(WorkpaperCell("B7", "Estimated Tax", tax, formula="=B5*B6",
                              source=AuthoritySource.WORKBOOK_FORMULA))
    wp.set_cell(WorkpaperCell("B8", "Net After Tax", net, formula="=B5-B7",
                              source=AuthoritySource.WORKBOOK_FORMULA))

    wp.notes.append("Prepared from current-year trial balance (synthetic).")
    wp.notes.append("Tax rate per management instruction memo (synthetic).")
    return wp


def make_defective_workpaper(seed: int = 20240101) -> Workpaper:
    """Build a workpaper seeded with realistic, catchable defects."""
    rng = random.Random(seed)
    wp, rev_a, rev_b, rev_c = _new_workpaper(rng)

    total_rev = rev_a + rev_b + rev_c
    tax_rate = 0.21
    correct_tax = round(total_rev * tax_rate, 2)

    # Defect 1: tie-out break -- stated total does not equal B2+B3+B4.
    broken_total = total_rev - 1_000.0

    # Defect 2: hard-coded value where a formula is expected (B7 has no formula).
    wrong_tax = round(correct_tax * 0.90, 2)  # understated, no formula backing

    # Defect 3: net uses the broken total, compounding the error.
    net = round(broken_total - wrong_tax, 2)

    wp.set_cell(WorkpaperCell("B2", "Revenue - Stream A", rev_a,
                              source=AuthoritySource.CURRENT_YEAR_SOURCE))
    wp.set_cell(WorkpaperCell("B3", "Revenue - Stream B", rev_b,
                              source=AuthoritySource.CURRENT_YEAR_SOURCE))
    wp.set_cell(WorkpaperCell("B4", "Revenue - Stream C", rev_c,
                              source=AuthoritySource.CURRENT_YEAR_SOURCE))
    # Stated total claims a formula but the value does not tie out to it.
    wp.set_cell(WorkpaperCell("B5", "Total Revenue", broken_total, formula="=B2+B3+B4",
                              source=AuthoritySource.WORKBOOK_FORMULA))
    # Defect 4: AI-assumption tax rate with no supporting authority.
    wp.set_cell(WorkpaperCell("B6", "Tax Rate (AI guess)", 0.18,
                              source=AuthoritySource.AI_ASSUMPTION))
    # Hard-coded tax, no formula -> Reviewer flags missing formula + mismatch.
    wp.set_cell(WorkpaperCell("B7", "Estimated Tax", wrong_tax, formula=None,
                              source=AuthoritySource.AI_ASSUMPTION))
    wp.set_cell(WorkpaperCell("B8", "Net After Tax", net, formula="=B5-B7",
                              source=AuthoritySource.WORKBOOK_FORMULA))

    # Defect 5: leaked internal/process language in a client-facing note.
    wp.notes.append(rng.choice(_LEAKED_PROCESS_PHRASES))
    wp.notes.append("Prepared in a hurry; needs a second look (synthetic).")
    return wp


def build_sample(kind: str = "defective", seed: int = 20240101) -> Workpaper:
    """Factory used by the CLI/tests. ``kind`` is ``'clean'`` or ``'defective'``."""
    if kind == "clean":
        return make_clean_workpaper(seed)
    if kind == "defective":
        return make_defective_workpaper(seed)
    raise ValueError(f"Unknown sample kind: {kind!r} (use 'clean' or 'defective')")


def write_xlsx(wp: Workpaper, path: str) -> str:
    """Write the workpaper to a tiny ``.xlsx`` file and return the path.

    Kept optional and dependency-light: only :mod:`openpyxl` (already
    installed) is used. The pipeline itself never requires Excel -- this is for
    a self-contained, Excel-shaped artifact.
    """
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Workpaper"
    sheet.append(["Ref", "Label", "Value", "Formula", "Source"])
    for cell in wp.ordered_cells():
        sheet.append([cell.ref, cell.label, cell.value, cell.formula or "",
                      cell.source.label])
    sheet.append([])
    sheet.append(["Notes"])
    for note in wp.notes:
        sheet.append(["", note])
    book.save(path)
    return path


if __name__ == "__main__":  # pragma: no cover - manual smoke test
    for sample_kind in ("clean", "defective"):
        sample = build_sample(sample_kind)
        print(f"=== {sample_kind} sample: {sample.entity} ({sample.period}) ===")
        for c in sample.ordered_cells():
            print(f"  {c.ref}: {c.label:<28} value={c.value} formula={c.formula}")
        print()
